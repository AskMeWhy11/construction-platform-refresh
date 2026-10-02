"""Экспорт CostRate в Excel (pivot: строки=города, колонки=классы).

Данные берутся тем же запросом и в том же порядке, что и на странице
«Себестоимость (таблица)» → «Таблица» (см. views_admin._cost_pivot_dataset).
"""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

CITY_HEADER = "ГОРОД"
SHEET_TITLE = "Себестоимость"
FILENAME_TEMPLATE = "costrate_pivot_{date}.xlsx"
MIN_COLUMN_WIDTH = 12
COLUMN_WIDTH_PADDING = 2


def build_pivot_workbook(dataset: dict) -> bytes:
    """
    dataset — результат _cost_pivot_dataset(): {"classes": [...], "rows": [...]}.

    Заголовок: «ГОРОД» + названия классов в порядке страницы.
    Строки: города в порядке страницы, значения — числами, NULL → пустая ячейка.
    """
    classes = dataset.get("classes") or []
    rows = dataset.get("rows") or []

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE

    bold = Font(bold=True)
    ws.append([CITY_HEADER] + [cls.name for cls in classes])
    for cell in ws[1]:
        cell.font = bold
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        values = [row.get("city_name") or ""]
        for cell in row.get("cells") or []:
            value = cell.get("value")
            # None → пустая ячейка; Decimal пишется числом
            values.append("" if value is None else value)
        ws.append(values)

    widths = [len(CITY_HEADER)]
    for row in rows:
        widths[0] = max(widths[0], len(row.get("city_name") or ""))
    for idx, cls in enumerate(classes, start=2):
        width = max(len(cls.name), len("0" * 12))
        widths.append(width)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = (
            max(MIN_COLUMN_WIDTH, width + COLUMN_WIDTH_PADDING)
        )
    ws.freeze_panes = "B2"

    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def export_filename(today: date | None = None) -> str:
    """costrate_pivot_YYYY-MM-DD.xlsx"""
    return FILENAME_TEMPLATE.format(date=(today or date.today()).isoformat())
