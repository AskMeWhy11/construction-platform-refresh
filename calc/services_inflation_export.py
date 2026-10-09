"""Выгрузка MonthlyInflation в Excel (колонка A — месяц, колонка B — инфляция)."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from .models import MonthlyInflation

SHEET_TITLE = "Инфляция"
HEADER_MONTH = "Месяц"
HEADER_RATE = "Инфляция"
# Тот же формат, что и у поля модели: столько знаков, сколько decimal_places.
RATE_NUMBER_FORMAT = "0." + "0" * MonthlyInflation._meta.get_field("rate").decimal_places

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_filename(now: datetime | None = None) -> str:
    """Имя файла только из ASCII — безопасно для Content-Disposition."""
    stamp = (now or datetime.now()).strftime("%Y%m%d")
    return f"monthly_inflation_{stamp}.xlsx"


def build_export_workbook() -> bytes:
    """
    Файл выгрузки: строка-шапка + по строке на месяц (сортировка по месяцу).
    Пустая таблица → файл только с шапкой.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE

    ws.append([HEADER_MONTH, HEADER_RATE])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14

    for month, rate in MonthlyInflation.objects.order_by("month").values_list("month", "rate"):
        ws.append([month, rate])
        ws.cell(row=ws.max_row, column=2).number_format = RATE_NUMBER_FORMAT

    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
