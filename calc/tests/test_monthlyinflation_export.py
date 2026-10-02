"""Тесты экспорта Excel для справочника «Инфляция (помесячно)»."""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Permission, User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from calc.models import MonthlyInflation
from calc.services_inflation_export import CONTENT_TYPE
from calc.services_inflation_import import import_from_bytes

EXPORT_URL_NAME = "admin:calc_monthlyinflation_export"


def read_rows(payload: bytes) -> list[tuple]:
    wb = load_workbook(BytesIO(payload), data_only=True, read_only=True)
    try:
        return [tuple(row) for row in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


class MonthlyInflationExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("boss", "boss@example.com", "pw")
        self.client.force_login(self.user)
        # Миграция-сид 0009 заполняет ~128 месяцев: для арифметики тестов чистим таблицу.
        MonthlyInflation.objects.all().delete()

    def test_changelist_has_export_button(self):
        response = self.client.get(reverse("admin:calc_monthlyinflation_changelist"))
        self.assertContains(response, "Экспорт Excel")
        self.assertContains(response, reverse(EXPORT_URL_NAME))

    def test_exports_xlsx_sorted_by_month(self):
        MonthlyInflation.objects.create(month=3, rate=Decimal("0.0762"))
        MonthlyInflation.objects.create(month=1, rate=Decimal("0.0732"))

        response = self.client.get(reverse(EXPORT_URL_NAME))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], CONTENT_TYPE)
        disposition = response["Content-Disposition"]
        self.assertIn("attachment;", disposition)
        self.assertIn("monthly_inflation_", disposition)
        self.assertTrue(disposition.rstrip('"').endswith(".xlsx"))

        rows = read_rows(response.content)
        self.assertEqual(rows[0], ("Месяц", "Инфляция"))
        self.assertEqual([r[0] for r in rows[1:]], [1, 3])
        self.assertEqual(Decimal(str(rows[1][1])), Decimal("0.0732"))
        self.assertEqual(Decimal(str(rows[2][1])), Decimal("0.0762"))

    def test_empty_table_exports_header_only(self):
        response = self.client.get(reverse(EXPORT_URL_NAME))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(read_rows(response.content), [("Месяц", "Инфляция")])

    def test_roundtrip_export_then_import(self):
        MonthlyInflation.objects.create(month=1, rate=Decimal("0.0732"))
        MonthlyInflation.objects.create(month=2, rate=Decimal("0.0747"))

        payload = self.client.get(reverse(EXPORT_URL_NAME)).content
        MonthlyInflation.objects.all().delete()

        created, updated = import_from_bytes(payload)

        self.assertEqual((created, updated), (2, 0))
        self.assertEqual(MonthlyInflation.objects.get(month=1).rate, Decimal("0.0732"))
        self.assertEqual(MonthlyInflation.objects.get(month=2).rate, Decimal("0.0747"))

    def test_view_permission_is_enough(self):
        user = User.objects.create_user("viewer", "v@example.com", "pw", is_staff=True)
        user.user_permissions.add(Permission.objects.get(codename="view_monthlyinflation"))
        self.client.force_login(user)

        response = self.client.get(reverse(EXPORT_URL_NAME))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], CONTENT_TYPE)

    def test_staff_without_permission_is_redirected(self):
        user = User.objects.create_user("nobody", "n@example.com", "pw", is_staff=True)
        self.client.force_login(user)

        response = self.client.get(reverse(EXPORT_URL_NAME))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], reverse("admin:calc_monthlyinflation_changelist"))
        self.assertNotIn("attachment", response.get("Content-Disposition", ""))

    def test_anonymous_redirected_to_login(self):
        self.client.logout()
        response = self.client.get(reverse(EXPORT_URL_NAME))
        self.assertEqual(response.status_code, 302)
        self.assertIn("/admin/login/", response["Location"])
