"""Тесты точности ставки инфляции (7 знаков после запятой)."""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from calc.models import MonthlyInflation
from calc.services_inflation_export import RATE_NUMBER_FORMAT
from calc.services_inflation_import import (
    MAX_RATE_DECIMALS,
    MAX_RATE_DIGITS,
    MAX_RATE_INT_DIGITS,
    import_from_bytes,
)

IMPORT_URL_NAME = "admin:calc_monthlyinflation_import"
EXPORT_URL_NAME = "admin:calc_monthlyinflation_export"


def make_xlsx(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


class MonthlyInflationPrecisionTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("boss", "boss@example.com", "pw")
        self.client.force_login(self.user)
        MonthlyInflation.objects.all().delete()

    def test_model_field_allows_seven_decimals(self):
        field = MonthlyInflation._meta.get_field("rate")
        self.assertEqual(field.max_digits, 9)
        self.assertEqual(field.decimal_places, 7)

    def test_limits_follow_field_meta(self):
        self.assertEqual((MAX_RATE_DIGITS, MAX_RATE_DECIMALS), (9, 7))
        # Целая часть не расширялась: как было 2 знака, так и осталось.
        self.assertEqual(MAX_RATE_INT_DIGITS, 2)

    def test_export_format_has_seven_decimals(self):
        self.assertEqual(RATE_NUMBER_FORMAT, "0.0000000")

    def test_import_keeps_seven_decimals(self):
        payload = make_xlsx([(1, "0,0732001"), (2, 0.1234567)])
        created, updated = import_from_bytes(payload)

        self.assertEqual((created, updated), (2, 0))
        self.assertEqual(MonthlyInflation.objects.get(month=1).rate, Decimal("0.0732001"))
        self.assertEqual(
            MonthlyInflation.objects.get(month=2).rate, Decimal("0.1234567")
        )

    def test_roundtrip_through_admin_export_and_import(self):
        MonthlyInflation.objects.create(month=1, rate=Decimal("0.0732001"))
        MonthlyInflation.objects.create(month=2, rate=Decimal("0.0800000"))

        response = self.client.get(reverse(EXPORT_URL_NAME))
        MonthlyInflation.objects.all().delete()

        created, updated = import_from_bytes(response.content)

        self.assertEqual((created, updated), (2, 0))
        self.assertEqual(MonthlyInflation.objects.get(month=1).rate, Decimal("0.0732001"))
        self.assertEqual(MonthlyInflation.objects.get(month=2).rate, Decimal("0.0800000"))

    def test_export_writes_full_precision_number(self):
        MonthlyInflation.objects.create(month=1, rate=Decimal("0.0732001"))

        response = self.client.get(reverse(EXPORT_URL_NAME))
        wb = load_workbook(BytesIO(response.content), data_only=True, read_only=True)
        try:
            cells = list(wb.active.iter_rows(values_only=True))
        finally:
            wb.close()

        self.assertEqual(cells[0], ("Месяц", "Инфляция"))
        self.assertEqual(Decimal(str(cells[1][1])), Decimal("0.0732001"))

    def test_export_number_format_row_is_seven_decimals(self):
        MonthlyInflation.objects.create(month=1, rate=Decimal("0.0732001"))

        response = self.client.get(reverse(EXPORT_URL_NAME))
        wb = load_workbook(BytesIO(response.content))
        try:
            self.assertEqual(wb.active.cell(row=2, column=2).number_format, "0.0000000")
        finally:
            wb.close()

    def test_import_rejects_eight_decimals(self):
        payload = make_xlsx([(1, "0.07320011")])
        response = self.client.post(
            reverse(IMPORT_URL_NAME),
            {"file": SimpleUploadedFile("x.xlsx", payload)},
            follow=True,
        )

        self.assertEqual(MonthlyInflation.objects.count(), 0)
        joined = " ".join(str(m) for m in response.context["messages"])
        self.assertIn("Строка 2", joined)
        self.assertIn("не помещается в формат поля", joined)

    def test_import_rejects_three_integer_digits(self):
        payload = make_xlsx([(1, "123.4567890")])
        response = self.client.post(
            reverse(IMPORT_URL_NAME),
            {"file": SimpleUploadedFile("x.xlsx", payload)},
            follow=True,
        )

        self.assertEqual(MonthlyInflation.objects.count(), 0)
        self.assertIn("не помещается в формат поля", " ".join(str(m) for m in response.context["messages"]))

    def test_import_accepts_max_boundary_value(self):
        created, _updated = import_from_bytes(make_xlsx([(1, "99.9999999")]))
        self.assertEqual(created, 1)
        self.assertEqual(MonthlyInflation.objects.get(month=1).rate, Decimal("99.9999999"))

    def test_legacy_four_decimal_rows_still_fine(self):
        MonthlyInflation.objects.create(month=1, rate=Decimal("0.0732"))
        created, updated = import_from_bytes(make_xlsx([(1, "0.0732")]))
        self.assertEqual((created, updated), (0, 1))
        self.assertEqual(MonthlyInflation.objects.get(month=1).rate, Decimal("0.0732"))
