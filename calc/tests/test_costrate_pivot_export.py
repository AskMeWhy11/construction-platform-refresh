"""Тесты экспорта Excel для страницы «Себестоимость (таблица)» → «Таблица»."""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from calc.models import BuildingClass, City, CostRate

EXPORT_URL_NAME = "admin:calc_costrate_pivot_export"
PIVOT_URL_NAME = "admin:calc_costrate_pivot"


def read_workbook(response):
    wb = load_workbook(BytesIO(response.content), data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return ws, rows


class CostRatePivotExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("boss", "boss@example.com", "pw")
        self.client.force_login(self.user)
        # Миграции-сиды 0004/0007 заполняют справочники: для проверки порядка и
        # пустых ячеек работаем на чистой таблице.
        CostRate.objects.all().delete()
        City.objects.all().delete()
        BuildingClass.objects.all().delete()

        self.cls_business = BuildingClass.objects.create(name="БИЗНЕС", order=2)
        self.cls_premium = BuildingClass.objects.create(name="ПРЕМИУМ", order=1)
        # порядок на странице — по (order, name): ПРЕМИУМ, БИЗНЕС
        self.classes = [self.cls_premium, self.cls_business]

        self.city_b = City.objects.create(name="Брянская область")
        self.city_a = City.objects.create(name="Алтайский край")
        self.cities = [self.city_a, self.city_b]  # порядок на странице — по name

        CostRate.objects.create(
            city=self.city_a, building_class=self.cls_premium,
            price_per_sqm=Decimal("112512.00"),
        )
        CostRate.objects.create(
            city=self.city_b, building_class=self.cls_business,
            price_per_sqm=Decimal("59673.00"),
        )

    def get_export(self):
        return self.client.get(reverse(EXPORT_URL_NAME))

    def test_page_has_export_button(self):
        response = self.client.get(reverse(PIVOT_URL_NAME))
        self.assertContains(response, "Экспорт Excel")
        self.assertContains(response, reverse(EXPORT_URL_NAME))

    def test_city_order_is_alphabetical_on_page(self):
        response = self.client.get(reverse(PIVOT_URL_NAME))
        names = [r["city_name"] for r in response.context["rows"]]
        self.assertEqual(names, [c.name for c in self.cities])

    def test_download_headers(self):
        response = self.get_export()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("costrate_pivot_", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])

    def test_header_row_order_matches_page(self):
        _, rows = read_workbook(self.get_export())
        self.assertEqual(rows[0], ["ГОРОД"] + [c.name for c in self.classes])

    def test_row_order_matches_page(self):
        _, rows = read_workbook(self.get_export())
        self.assertEqual([r[0] for r in rows[1:]], [c.name for c in self.cities])

    def test_values_are_numbers_and_missing_cells_are_empty(self):
        _, rows = read_workbook(self.get_export())

        # колонки: ГОРОД, ПРЕМИУМ, БИЗНЕС
        a_row, b_row = rows[1], rows[2]
        self.assertEqual(a_row[0], "Алтайский край")
        self.assertEqual(a_row[1], 112512)
        self.assertIsNone(a_row[2])
        self.assertEqual(b_row[0], "Брянская область")
        self.assertIsNone(b_row[1])
        self.assertEqual(b_row[2], 59673)

    def test_empty_table_keeps_header_only(self):
        CostRate.objects.all().delete()
        City.objects.all().delete()

        _, rows = read_workbook(self.get_export())
        self.assertEqual(rows, [["ГОРОД"] + [c.name for c in self.classes]])

    def test_anonymous_redirected_to_login(self):
        self.client.logout()
        response = self.get_export()
        self.assertEqual(response.status_code, 302)
        self.assertIn("/admin/login/", response["Location"])

    def test_non_staff_redirected_to_login(self):
        user = User.objects.create_user("viewer", "v@example.com", "pw")
        self.client.force_login(user)

        response = self.get_export()
        self.assertEqual(response.status_code, 302)
        self.assertIn("/admin/login/", response["Location"])
